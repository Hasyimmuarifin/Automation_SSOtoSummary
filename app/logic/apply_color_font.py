from openpyxl.styles import Font
# from openpyxl.utils import get_column_letter

def apply_status_font(sheet, status_col_idx=69, start_col=14, end_col=62):
    """
    Ubah warna font pada kolom blendingan (N sampai BJ) berdasarkan status (kolom BQ).
    - Status == "Plan"        → font biru
    - Status == "Completed" / "Loading" / "In Progress" → font Verdana hitam size 9
    """
    max_row = sheet.max_row

    for row in range(2, max_row + 1):  # skip header row (mulai dari baris 2)
        status_val = sheet.cell(row=row, column=status_col_idx).value
        if not status_val:
            continue

        # tentukan warna font
        if str(status_val).strip().lower() == "plan":
            font_color = "0070C0"  # biru
        elif str(status_val).strip().lower() in ["completed", "loading", "in progress"]:
            font_color = "000000"  # hitam
        else:
            continue  # status tidak dikenali → biarkan default

        # loop kolom N (14) sampai BJ (62)
        for col in range(start_col, end_col + 1):
            cell = sheet.cell(row=row, column=col)
            if cell.value is not None:
                cell.font = Font(name="Verdana", size=9, color=font_color)

    print("🎨 Font status berhasil diterapkan (Plan=biru, Completed/Loading/In Progress=hitam)")