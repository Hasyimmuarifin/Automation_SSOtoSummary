# backup_restore_plan.py

import openpyxl
from openpyxl.utils import get_column_letter

def backup_plan_rows(wb, sheet_b, backup_sheet_name="Backup_Plan", debug=True):
    """
    Backup baris dengan status 'Plan' ke sheet sementara.
    - Jika 'Plan' → backup Month, Company, Vessel, End User, serta AKC–AKQ (numeric).
    - debug=True → cetak informasi tambahan untuk verifikasi AON/AOP/AOR/AOT
    """
    # hapus sheet backup jika sudah ada
    if backup_sheet_name in wb.sheetnames:
        del wb[backup_sheet_name]
    ws_backup = wb.create_sheet(backup_sheet_name)

    # header utama (data numeric)
    headers = (
        ['Month', 'Company', 'Vessel', 'End User'] + [f"AK{chr(c)}" for c in range(ord('C'), ord('R'))]  # AKC–AKQ 
        + ["AON", "AOP", "AOR", "AOT"]  # kolom tambahan
    )
    ws_backup.append(headers)

    # kolom index
    COL_STATUS = 69   # BQ
    COL_C = 3         # Month
    COL_D = 4         # Company
    COL_E = 5         # Vessel
    COL_G = 7         # End User
    COL_AKC = 965     # AKC
    COL_AKQ = 979     # AKQ
    COL_AON = 1080
    COL_AOP = 1082
    COL_AOR = 1084
    COL_AOT = 1086

    # sanity checks (debug)
    if debug:
        print(f"[DEBUG] sheet_b max_row={sheet_b.max_row}, max_column={sheet_b.max_column}")
        print(f"[DEBUG] Indeks penting: AKC={COL_AKC}..AKQ={COL_AKQ} (count={COL_AKQ - COL_AKC + 1}), AON={COL_AON}, AOP={COL_AOP}, AOR={COL_AOR}, AOT={COL_AOT}, COL_MAX={COL_MAX}")
        if COL_MAX < COL_AOT:
            print(f"[WARN] COL_MAX ({COL_MAX}) < COL_AOT ({COL_AOT}) → akan menyesuaikan COL_MAX = COL_AOT")
            COL_MAX = COL_AOT
        # warn jika indeks kolom melebihi sheet actual
        if any(c > sheet_b.max_column for c in (COL_AKC, COL_AKQ, COL_AON, COL_AOP, COL_AOR, COL_AOT, COL_MAX)):
            print("[WARN] Salah satu indeks kolom yang dipakai melebihi sheet_b.max_column. Periksa mapping kolom Anda.")

    # hitungan expected
    ak_count = COL_AKQ - COL_AKC + 1
    expected_len = 4 + ak_count + 4  # Month,Company,Vessel,EndUser + AKC..AKQ + AON..AOT

    for row in range(2, sheet_b.max_row + 1):
        status = sheet_b.cell(row=row, column=COL_STATUS).value
        if status == "Plan":
            month = sheet_b.cell(row=row, column=COL_C).value
            company = sheet_b.cell(row=row, column=COL_D).value
            vessel = sheet_b.cell(row=row, column=COL_E).value
            enduser = sheet_b.cell(row=row, column=COL_G).value

            if debug:
                print(f"\n[DEBUG] ---- Processing row {row} ----")
                print(f"[DEBUG] Status='{status}' Month={month!r}, Company={company!r}, Vessel={vessel!r}, EndUser={enduser!r}")

            # ambil nilai numeric AKC–AKQ
            values = []
            for col in range(COL_AKC, COL_AKQ + 1):
                cell = sheet_b.cell(row=row, column=col).value
                values.append(cell if isinstance(cell, (int, float)) else None)
            if debug:
                first_col_letter = get_column_letter(COL_AKC)
                last_col_letter = get_column_letter(COL_AKQ)
                print(f"[DEBUG] AK values ({first_col_letter}{row}..{last_col_letter}{row}) count={len(values)} sample: {values[:3]} ... {values[-3:]}")

            # ambil tambahan AON, AOP, AOR, AOT
            aon_values = []
            for col in (COL_AON, COL_AOP, COL_AOR, COL_AOT):
                val = sheet_b.cell(row=row, column=col).value
                values.append(val)   # cukup append val
                aon_values.append((col, get_column_letter(col), val))
            if debug:
                for col_idx, col_letter, val in aon_values:
                    print(f"[DEBUG] AON-block: Col {col_idx} ({col_letter}{row}) = {val!r}")
                # check if all None
                if all(v is None for (_, _, v) in aon_values):
                    print(f"[WARN] Semua nilai AON..AOT None pada row {row} (cek mapping kolom / apakah cell berisi formula tanpa nilai).")

            row_data = [month, company, vessel, enduser] + values

            # debug length check sebelum append
            if debug:
                print(f"[DEBUG] row_data length={len(row_data)} expected={expected_len}. row_data head: {row_data[:8]}")

            ws_backup.append(row_data)

            # format kolom Month (kolom A pada sheet backup)
            ws_backup.cell(row=ws_backup.max_row, column=1).number_format = "mmm"
            print(f"[BACKUP] Row {row} → appended to Backup_Plan row {ws_backup.max_row}")

def restore_plan_rows(wb, sheet_b, backup_sheet_name="Backup_Plan"):
    """
    Restore data dari sheet Backup_Plan ke sheet ITM Summary (sheet_b).
    Pencocokan berdasarkan 4 kolom: Month, Company, Vessel, End User.
    Jika cocok, isi kembali nilai AKC–AKQ.
    - Restore numeric AKC–AKQ berdasarkan 4 kolom kunci.
    Setelah restore selesai, sheet Backup_Plan dihapus.
    """

    if backup_sheet_name not in wb.sheetnames:
        print("[RESTORE] Tidak ada sheet Backup_Plan. Restore dibatalkan.")
        return

    ws_backup = wb[backup_sheet_name]

    # index kolom pada sheet ITM Summary
    COL_C = 3         # Month
    COL_D = 4         # Company
    COL_E = 5         # Vessel
    COL_G = 7         # End User
    COL_AKC = 965     # AKC
    COL_AKQ = 979     # AKQ
    COL_AON = 1080
    COL_AOP = 1082
    COL_AOR = 1084
    COL_AOT = 1086

    # iterasi semua data di backup (mulai dari row 2, karena row 1 adalah header)
    for row in range(2, ws_backup.max_row + 1):
        month_bkp = ws_backup.cell(row=row, column=1).value
        company_bkp = ws_backup.cell(row=row, column=2).value
        vessel_bkp = ws_backup.cell(row=row, column=3).value
        enduser_bkp = ws_backup.cell(row=row, column=4).value

        values_bkp = [
            ws_backup.cell(row=row, column=col).value
            for col in range(5, ws_backup.max_column + 1)
        ]

        # cari baris yang cocok di sheet ITM Summary
        for r in range(2, sheet_b.max_row + 1):
            month_val = sheet_b.cell(row=r, column=COL_C).value
            company_val = sheet_b.cell(row=r, column=COL_D).value
            vessel_val = sheet_b.cell(row=r, column=COL_E).value
            enduser_val = sheet_b.cell(row=r, column=COL_G).value

            if (
                month_val == month_bkp
                and company_val == company_bkp
                and vessel_val == vessel_bkp
                and enduser_val == enduser_bkp
            ):
                # cocok → restore nilai numeric AKC–AKQ
                for idx, col in enumerate(range(COL_AKC, COL_AKQ + 1)):
                    val = values_bkp[idx] if idx < len(values_bkp) else None
                    if val is not None:
                        sheet_b.cell(row=r, column=col).value = val

                # restore tambahan AON–AOT (4 kolom setelah AKQ)
                extra_cols = [COL_AON, COL_AOP, COL_AOR, COL_AOT]
                for i, col in enumerate(extra_cols, start=(COL_AKQ - COL_AKC + 1)):
                    val = values_bkp[i] if i < len(values_bkp) else None
                    if val is not None:
                        sheet_b.cell(row=r, column=col).value = val

                print(f"[RESTORE] Row {r} diperbarui dari backup (Month={month_bkp}, Company={company_bkp})")
                break  # sudah ketemu baris, tidak perlu cari 

    # hapus sheet backup setelah selesai
    del wb[backup_sheet_name]
    print("[RESTORE] Sheet Backup_Plan berhasil dihapus setelah restore.")

def clear_aon_block(sheet_b, debug=True):
    """
    Hapus value di kolom AON, AOP, AOR, AOT untuk baris dengan Status='Plan' (kolom BQ).
    Dipanggil setelah proses perbulan selesai dan sebelum restore_plan_rows().
    """
    COL_STATUS = 69    # BQ
    COL_AON = 1080
    COL_AOP = 1082
    COL_AOR = 1084
    COL_AOT = 1086

    cleared_rows = 0

    for row in range(2, sheet_b.max_row + 1):
        status = sheet_b.cell(row=row, column=COL_STATUS).value
        if status == "Plan":
            for col in (COL_AON, COL_AOP, COL_AOR, COL_AOT):
                if sheet_b.cell(row=row, column=col).value is not None:
                    sheet_b.cell(row=row, column=col).value = None
                    if debug:
                        from openpyxl.utils import get_column_letter
                        print(f"[CLEAR] Row {row}, Col {get_column_letter(col)} → cleared")
            cleared_rows += 1

    print(f"[CLEAR] Total {cleared_rows} rows dengan Status='Plan' dibersihkan")